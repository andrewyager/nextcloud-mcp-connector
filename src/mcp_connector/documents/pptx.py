"""PowerPoint to Markdown: one section per slide, in slide order. Images are dropped."""

import io

from .docx import _cell, _table_lines

__all__ = ["to_markdown"]


def to_markdown(data: bytes) -> str:
    import pptx
    from pptx.shapes.autoshape import Shape
    from pptx.shapes.graphfrm import GraphicFrame

    presentation = pptx.Presentation(io.BytesIO(data))
    out: list[str] = []
    for number, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape is not None else ""
        out.append(f"## Slide {number}: {title}" if title else f"## Slide {number}")
        out.append("")
        for shape in slide.shapes:
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            if isinstance(shape, Shape) and shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in paragraph.runs).strip()
                    if text:
                        out.append(text)
                if shape.text_frame.text.strip():
                    out.append("")
            elif isinstance(shape, GraphicFrame) and shape.has_table:
                rows = [[_cell(cell.text) for cell in row.cells] for row in shape.table.rows]
                out.extend(_table_lines(rows))
                out.append("")
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                out.append(f"Notes: {' '.join(notes.text.split())}")
                out.append("")
    return "\n".join(out).strip() + "\n"
