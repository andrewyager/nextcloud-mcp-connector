"""Excel to Markdown: one heading per sheet, a pipe table of stored values, capped rows."""

import io
from pathlib import Path

import openpyxl
import pytest

from mcp_connector.documents import xlsx as xlsx_conv
from mcp_connector.errors import REASON_GUARD_TRIPPED, ToolError

FIXTURE = Path(__file__).resolve().parents[1] / "fixtures" / "documents" / "sample.xlsx"


def test_each_sheet_becomes_a_heading_and_a_pipe_table() -> None:
    text = xlsx_conv.to_markdown(FIXTURE.read_bytes())

    assert "## Report" in text
    assert "## Appendix" in text
    assert "| Item | Amount |" in text
    assert "| Servers | 12 |" in text
    assert "| Date | 2026-09-30 |" in text
    assert text.index("## Report") < text.index("## Appendix")


def test_trailing_empty_rows_and_columns_are_trimmed() -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = "a"
    sheet["B1"] = "b"
    sheet["A2"] = "c"
    sheet["E9"] = None
    buffer = io.BytesIO()
    workbook.save(buffer)

    text = xlsx_conv.to_markdown(buffer.getvalue())
    lines = [line for line in text.splitlines() if line.startswith("|")]
    assert lines == ["| a | b |", "| --- | --- |", "| c |  |"]


def test_rows_above_the_cap_are_cut_with_a_note(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(xlsx_conv, "MAX_ROWS", 3)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    for i in range(10):
        sheet.append([i])
    buffer = io.BytesIO()
    workbook.save(buffer)

    text = xlsx_conv.to_markdown(buffer.getvalue())
    assert text.count("\n| ") <= 4  # header row, separator, two more rows at most
    assert "rows omitted" in text


def test_too_many_sheets_are_refused(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(xlsx_conv, "MAX_SHEETS", 2)
    workbook = openpyxl.Workbook()
    workbook.create_sheet("Two")
    workbook.create_sheet("Three")
    buffer = io.BytesIO()
    workbook.save(buffer)

    with pytest.raises(ToolError) as info:
        xlsx_conv.to_markdown(buffer.getvalue())
    assert info.value.reason == REASON_GUARD_TRIPPED
